"""Leitura estrutural de arquivos CSV e XLSX (seção 10, passos 1 a 4 do pipeline)."""

import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from enum import StrEnum

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 50_000
MAX_SHEETS = 20
SUPPORTED_EXTENSIONS = (".csv", ".xlsx")

# Chave reservada na linha bruta guardada em staging para os tipos nativos do Excel.
NATIVE_TYPES_KEY = "__native_types__"


class NativeType(StrEnum):
    """Células que o Excel já entrega tipadas; não dependem da convenção escolhida."""

    NUMBER = "number"  # texto canônico com ponto decimal, ex.: 1234.56
    DATE = "date"  # texto ISO, ex.: 2026-09-30


class FileRejected(Exception):
    """Arquivo recusado antes de qualquer leitura de linhas."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ParsedRow:
    row_number: int  # linha no arquivo, contando o cabeçalho como linha 1
    values: dict[str, str]
    native_types: dict[str, NativeType] = field(default_factory=dict)

    def to_staging(self) -> dict:
        if not self.native_types:
            return dict(self.values)
        return {**self.values, NATIVE_TYPES_KEY: dict(self.native_types)}

    @classmethod
    def from_staging(cls, row_number: int, raw: dict) -> "ParsedRow":
        values = dict(raw)
        types = values.pop(NATIVE_TYPES_KEY, {}) or {}
        return cls(row_number, values, {col: NativeType(t) for col, t in types.items()})


@dataclass(frozen=True)
class ParsedFile:
    columns: list[str]
    rows: list[ParsedRow]
    sha256: str
    delimiter: str | None = None
    sheets: list[str] = field(default_factory=list)
    sheet: str | None = None


def _check_columns(header: list[str]) -> list[str]:
    columns = [col.strip() for col in header]
    if not any(columns):
        raise FileRejected(
            "MISSING_HEADER", "A primeira linha precisa conter os nomes das colunas."
        )
    if any(not col for col in columns):
        raise FileRejected("MISSING_HEADER", "Há colunas sem nome no cabeçalho.")
    if len(set(columns)) != len(columns):
        raise FileRejected("DUPLICATE_COLUMN", "O cabeçalho tem colunas com nomes repetidos.")
    if NATIVE_TYPES_KEY in columns:
        raise FileRejected("MISSING_HEADER", f"O nome de coluna '{NATIVE_TYPES_KEY}' é reservado.")
    return columns


def _check_row_limit(rows: list[ParsedRow]) -> None:
    if len(rows) >= MAX_ROWS:
        raise FileRejected("TOO_MANY_ROWS", f"O arquivo excede o limite de {MAX_ROWS} linhas.")


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileRejected("INVALID_ENCODING", "Não foi possível ler o arquivo como texto.")


def _parse_csv(content: bytes, sha256: str) -> ParsedFile:
    if b"\x00" in content[:4096]:
        raise FileRejected("UNSUPPORTED_FILE_TYPE", "O conteúdo não parece ser um CSV de texto.")

    text = _decode(content)
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        header = next(reader)
    except StopIteration:
        raise FileRejected("EMPTY_FILE", "O arquivo não tem cabeçalho.") from None
    while header and not header[-1].strip():  # separador sobrando no fim do cabeçalho
        header = header[:-1]
    columns = _check_columns(header)

    rows: list[ParsedRow] = []
    for values in reader:
        if not any(v.strip() for v in values):
            continue  # linha totalmente vazia
        _check_row_limit(rows)
        padded = values + [""] * (len(columns) - len(values))
        rows.append(ParsedRow(reader.line_num, dict(zip(columns, padded, strict=False))))

    return ParsedFile(columns=columns, rows=rows, sha256=sha256, delimiter=delimiter)


def _cell(value: object) -> tuple[str, NativeType | None]:
    """Converte uma célula do Excel em texto, marcando números e datas nativos."""
    if value is None:
        return "", None
    if isinstance(value, bool):
        return str(value).upper(), None
    if isinstance(value, datetime):
        return value.date().isoformat(), NativeType.DATE
    if isinstance(value, date):
        return value.isoformat(), NativeType.DATE
    if isinstance(value, int):
        return str(value), NativeType.NUMBER
    if isinstance(value, float):
        # repr dá a menor representação que reproduz o float (0.1 -> "0.1").
        return format(Decimal(repr(value)), "f"), NativeType.NUMBER
    return str(value).strip(), None


def _parse_xlsx(content: bytes, sha256: str, sheet: str | None) -> ParsedFile:
    # Importação adiada: só quem envia XLSX paga o custo de carregar o openpyxl.
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise FileRejected(
                    "FILE_TOO_LARGE", "O conteúdo descompactado do XLSX é grande demais."
                )
            names = archive.namelist()
    except zipfile.BadZipFile:
        raise FileRejected(
            "UNSUPPORTED_FILE_TYPE", "O conteúdo não é um arquivo XLSX válido."
        ) from None
    if "xl/vbaProject.bin" in names:
        raise FileRejected("MACROS_NOT_ALLOWED", "Arquivos com macros não são aceitos.")

    try:
        # read_only lê em streaming; data_only usa o último valor calculado das fórmulas
        # e nada é executado.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise FileRejected(
            "UNSUPPORTED_FILE_TYPE", "O conteúdo não é um arquivo XLSX válido."
        ) from None

    try:
        sheets = list(workbook.sheetnames)
        if len(sheets) > MAX_SHEETS:
            raise FileRejected(
                "TOO_MANY_SHEETS", f"O arquivo excede o limite de {MAX_SHEETS} abas."
            )
        chosen = sheet if sheet is not None else sheets[0]
        if chosen not in sheets:
            raise FileRejected("UNKNOWN_SHEET", f"A aba '{chosen}' não existe no arquivo.")

        lines = workbook[chosen].iter_rows(values_only=True)
        header = next(lines, None)
        if header is None:
            raise FileRejected("EMPTY_FILE", "A aba escolhida está vazia.")
        while header and _cell(header[-1])[0] == "":  # colunas vazias à direita
            header = header[:-1]
        columns = _check_columns([_cell(v)[0] for v in header])

        rows: list[ParsedRow] = []
        for line_number, cells in enumerate(lines, start=2):
            converted = [_cell(v) for v in cells[: len(columns)]]
            if not any(text for text, _ in converted):
                continue
            _check_row_limit(rows)
            converted += [("", None)] * (len(columns) - len(converted))
            values = {col: text for col, (text, _) in zip(columns, converted, strict=True)}
            native = {col: kind for col, (_, kind) in zip(columns, converted, strict=True) if kind}
            rows.append(ParsedRow(line_number, values, native))
    finally:
        workbook.close()

    return ParsedFile(columns=columns, rows=rows, sha256=sha256, sheets=sheets, sheet=chosen)


def parse_file(filename: str, content: bytes, sheet: str | None = None) -> ParsedFile:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if f".{extension}" not in SUPPORTED_EXTENSIONS:
        raise FileRejected(
            "UNSUPPORTED_FILE_TYPE", "Formato não suportado. Envie um arquivo .csv ou .xlsx."
        )
    if len(content) == 0:
        raise FileRejected("EMPTY_FILE", "O arquivo está vazio.")
    if len(content) > MAX_FILE_BYTES:
        raise FileRejected("FILE_TOO_LARGE", "O arquivo excede o limite de 5 MB.")

    sha256 = hashlib.sha256(content).hexdigest()
    if extension == "xlsx":
        return _parse_xlsx(content, sha256, sheet)
    return _parse_csv(content, sha256)
